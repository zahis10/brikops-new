"""Additional safe quantities sheet; legacy sheets remain unchanged."""
from openpyxl.styles import Font
from contractor_ops.xlsx_safe import append_row, set_cell

QUANTITY_HEADERS = ['מס׳ סעיף', 'תיאור', 'יח׳', 'מצטבר קודם', 'בוצע החודש', 'מצטבר',
                    'מחיר יח׳', 'סכום החודש', 'סכום מצטבר', 'דירות החודש', 'ראיות']


def add_quantities_sheet(wb, project, account, contractor):
    if not contractor.get('items'):
        return None
    # Lazy read-only imports avoid a cycle with the existing workbook builder.
    from contractor_ops.monthly_close_export import _sheet, _il_date, _month_label
    ws = _sheet(wb, 'כמויות לחשבון', QUANTITY_HEADERS, 4)
    ws.column_dimensions['J'].width, ws.column_dimensions['K'].width = 16, 40
    label = account.get('label') or _month_label(account.get('month'))
    set_cell(ws, 1, 1, f"כמויות לחשבון {contractor.get('account_no', '')} — {contractor.get('name', '')} — {label}")
    ws['A1'].font = Font(bold=True, size=14)
    month = account.get('month', '')
    period = f'{month[5:7]}/{month[:4]}'
    status = (f"נסגר {_il_date(account.get('closed_at'))} ע״י {(account.get('closed_by') or {}).get('name', '')}"
              if account.get('status') == 'closed' else 'טיוטה — החודש עדיין פתוח')
    set_cell(ws, 2, 1, f'חודש ביצוע {period} · {status}')
    money = contractor.get('totals_money')
    set_cell(ws, 3, 1, 'סכומים לפי המחירים שהוזנו ב-BrikOps · לפני מע״מ, לפני עכבון והצמדה'
             if money else 'כמויות בלבד — המחירים במערכת החשבונות')
    for line in contractor['items']:
        measured = line.get('source') == 'measured'
        evidence, entry = line.get('evidence_rows'), line.get('measurement')
        if measured:
            reference = (f"הוקלד {_il_date(entry.get('at'))} · {(entry.get('by') or {}).get('name', '')}"
                         + (f" · {entry['note']}" if entry.get('note') else '')) if entry else 'לא הוקלד החודש'
        else:
            reference = f'גיליון ראיות · שורות {evidence[0]}–{evidence[1]}' if evidence else '—'
        units = '—' if measured else str(line.get('units_this_month', 0))
        if not measured and line.get('corrections'):
            units += f" (−{line['corrections']})"
        numbers = [line.get(k) for k in ('prev_cumulative_qty', 'this_month_qty', 'cumulative_qty',
                                        'unit_price', 'this_month_amount', 'cumulative_amount')]
        append_row(ws, [line.get('code') or '—', line.get('description', ''), line.get('unit', ''),
                        *[n if n is not None else '' for n in numbers], units, reference])
    if money:
        append_row(ws, ['סה״כ', '', '', '', '', '', '', money['this_month'], money['cumulative'], '', ''])
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
    if any(l.get('corrections') for l in contractor['items']):
        append_row(ws, ['', 'תיקונים שנכללו בחודש זה — כבר מופחתים ב״בוצע החודש״ (גיליון תיקונים)', *[''] * 9])
    return ws
