"""Pure, safe per-contractor monthly progress workbook builder."""
from datetime import datetime
from io import BytesIO
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from contractor_ops.execution_matrix_export import HEADER_FILL
from contractor_ops.utils.timezone import IL_TZ
from contractor_ops.xlsx_safe import append_row, set_cell


SUMMARY_HEADERS = [
    'קבלן', 'שלב', 'בניין', 'סה״כ דירות', 'בוצע החודש', 'מצטבר',
    'אחוז', 'אושרו בבקרת ביצוע', 'סומנו ידנית',
]
EVIDENCE_HEADERS = [
    '#', 'בניין', 'קומה', 'דירה', 'שלב', 'תאריך ביצוע',
    'אושר ע״י', 'מקור', 'קישור',
]
CORRECTION_HEADERS = ['שלב', 'בניין', 'דירה', 'נספר בחודש', 'סיבה']
MONTHS = (
    'ינואר', 'פברואר', 'מרץ', 'אפריל', 'מאי', 'יוני',
    'יולי', 'אוגוסט', 'ספטמבר', 'אוקטובר', 'נובמבר', 'דצמבר',
)


def _il_date(value):
    """Return a display date; malformed or missing external dates are not trusted."""
    if not value:
        return '—'
    try:
        dt = datetime.fromisoformat(str(value).replace('Z', '+00:00'))
        if dt.tzinfo is None:
            from datetime import timezone
            dt = dt.replace(tzinfo=timezone.utc)
        dt = dt.astimezone(IL_TZ)
        return f'{dt.day}.{dt.month}.{dt.year}'
    except (ValueError, TypeError, OverflowError):
        return '—'


def _month_label(month):
    if month == 'baseline':
        return 'לפני החשבון הראשון ב-BrikOps'
    try:
        year, number = str(month).split('-')
        if len(year) == 4 and len(number) == 2 and 1 <= int(number) <= 12:
            return f'{MONTHS[int(number) - 1]} {year}'
    except (TypeError, ValueError, AttributeError):
        pass
    return '—'


def _sheet(wb, name, headers=None, header_row=1):
    ws = wb.create_sheet(name)
    ws.sheet_view.rightToLeft = True
    ws.freeze_panes = f'A{header_row + 1}'
    if headers:
        for index, value in enumerate(headers, 1):
            set_cell(ws, header_row, index, value)
        for cell in ws[header_row]:
            cell.fill = HEADER_FILL
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='right', wrap_text=True)
    for col in 'ABCDEFGHI':
        ws.column_dimensions[col].width = 22
    ws.column_dimensions['A'].width = 25
    ws.column_dimensions['D'].width = 14
    ws.column_dimensions['I'].width = 60 if name == 'ראיות' else 55
    return ws


def _link(base_url, project_id, row):
    root = (base_url or '').rstrip('/')
    pid = quote(str(project_id or ''), safe='')
    if row.get('source') == 'qc' and row.get('qc_run_id') and row.get('floor_id'):
        floor = quote(str(row['floor_id']), safe='')
        run = quote(str(row['qc_run_id']), safe='')
        stage = quote(str(row.get('stage_id') or ''), safe='')
        unit = quote(str(row.get('unit_id') or ''), safe='')
        return f'{root}/projects/{pid}/floors/{floor}/qc/{run}/stage/{stage}?unitId={unit}'
    return f'{root}/projects/{pid}/execution-matrix'


def build_monthly_close_xlsx(project, account, contractor, base_url):
    """Create a workbook from stored snapshot rows (or the same live account shape)."""
    wb = Workbook()
    wb.remove(wb.active)
    summary = _sheet(wb, 'סיכום', SUMMARY_HEADERS, 4)
    set_cell(summary, 1, 1, (
        f"דוח התקדמות לחשבון — {contractor.get('name', '')} — "
        f"{account.get('label', _month_label(account.get('month')))}"
    ))
    summary['A1'].font = Font(bold=True, size=14)
    if account.get('status') == 'closed':
        closer = account.get('closed_by') or {}
        set_cell(summary, 2, 1, f"נסגר {_il_date(account.get('closed_at'))} ע״י {closer.get('name', '')}")
    else:
        set_cell(summary, 2, 1, f"טיוטה — החודש עדיין פתוח (נכון ל-{_il_date(datetime.now(IL_TZ).isoformat())})")

    for stage in contractor.get('stages') or []:
        for building in stage.get('by_building') or []:
            append_row(summary, [
                contractor.get('name', ''), stage.get('title', ''),
                building.get('building_name', ''),
                building.get('total_units', 0), building.get('this_month', 0),
                building.get('cumulative', 0),
                round(100 * building.get('cumulative', 0) /
                      building['total_units']) if building.get('total_units') else 0,
                building.get('qc_count', 0), building.get('manual_count', 0),
            ])
    totals = contractor.get('totals') or {}
    append_row(summary, [
        'סה״כ', '', '', '',
        totals.get('this_month', 0), totals.get('cumulative', 0), '',
        totals.get('qc', 0), totals.get('manual', 0),
    ])
    for cell in summary[summary.max_row]:
        cell.font = Font(bold=True)

    for correction in contractor.get('corrections') or []:
        counted = correction.get('counted_in_month')
        origin = (
            '(נספרה לפני החשבון הראשון ב-BrikOps)'
            if counted == 'baseline' else f'(נספרה ב-{_month_label(counted)})'
        )
        reason = f"תיקון: דירה {correction.get('unit_no', '')} נפתחה מחדש {origin}"
        append_row(summary, [
            contractor.get('name', ''), correction.get('stage_title', ''),
            correction.get('building_name', ''), '', -1, '', '', '', reason,
        ])

    evidence = _sheet(wb, 'ראיות', EVIDENCE_HEADERS)
    row_number = 0
    for stage in contractor.get('stages') or []:
        for row in stage.get('evidence') or []:
            row_number += 1
            source = 'בקרת ביצוע' if row.get('source') == 'qc' else 'סימון ידני'
            if row.get('late_from_month'):
                source = (
                    'אושר בבקרת ביצוע' if row.get('source') == 'qc' else 'סומן ידנית'
                ) + f" · אחרי סגירת {_month_label(row['late_from_month'])}"
            append_row(evidence, [
                row_number, row.get('building_name', ''), row.get('floor_number', ''),
                row.get('unit_no', ''), stage.get('title', ''),
                _il_date(row.get('completed_at')), row.get('actor_name', ''),
                source, _link(base_url, project.get('id'), row),
            ])

    if contractor.get('corrections'):
        corrections = _sheet(wb, 'תיקונים', CORRECTION_HEADERS)
        for row in contractor['corrections']:
            append_row(corrections, [
                row.get('stage_title', ''), row.get('building_name', ''),
                row.get('unit_no', ''), _month_label(row.get('counted_in_month')),
                'נפתחה מחדש',
            ])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf