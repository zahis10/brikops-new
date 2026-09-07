from datetime import datetime
from io import BytesIO

from openpyxl import Workbook, load_workbook

from contractor_ops.xlsx_safe import append_row, set_cell


RISKY_VALUES = (
    '=1+1',
    '+1',
    '-1',
    '@SUM(A1)',
    '\t=x',
    "=cmd|' /C calc'!A0",
)


def _round_trip(workbook):
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return load_workbook(output)


def test_set_cell_stores_formula_like_values_as_verbatim_strings():
    workbook = Workbook()
    sheet = workbook.active
    for column, value in enumerate(RISKY_VALUES, start=1):
        set_cell(sheet, 1, column, value)

    loaded = _round_trip(workbook).active
    for column, value in enumerate(RISKY_VALUES, start=1):
        cell = loaded.cell(row=1, column=column)
        assert cell.data_type == 's'
        assert cell.value == value


def test_set_cell_preserves_plain_and_non_string_types():
    workbook = Workbook()
    sheet = workbook.active
    values = ('plain text', 42, None, datetime(2026, 9, 7, 12, 30))
    for column, value in enumerate(values, start=1):
        set_cell(sheet, 1, column, value)

    loaded = _round_trip(workbook).active
    cells = [loaded.cell(row=1, column=column) for column in range(1, 5)]
    assert [cell.value for cell in cells] == list(values)
    assert [cell.data_type for cell in cells] == ['s', 'n', 'n', 'd']


def test_append_row_neutralizes_only_formula_like_values():
    workbook = Workbook()
    sheet = workbook.active
    values = ['plain', '=1+1', 5, '@cmd']

    append_row(sheet, values)

    loaded = _round_trip(workbook).active
    assert [cell.value for cell in loaded[1]] == values
    assert [cell.data_type for cell in loaded[1]] == ['s', 's', 'n', 's']