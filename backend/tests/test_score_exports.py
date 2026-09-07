from io import BytesIO

from openpyxl import Workbook, load_workbook

from contractor_ops.safety.score_exports import _write_sheet


def test_write_sheet_neutralizes_formula_like_values():
    workbook = Workbook()
    workbook.remove(workbook.active)
    _write_sheet(
        workbook,
        'בדיקה',
        ['א', 'ב', 'ג'],
        [['=cmd', '-1', 'text']],
        [10, 10, 10],
    )
    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    sheet = load_workbook(output)['בדיקה']
    assert [cell.value for cell in sheet[2]] == ['=cmd', '-1', 'text']
    assert [cell.data_type for cell in sheet[2]] == ['s', 's', 's']