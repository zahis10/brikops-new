"""Contract item regression coverage, isolated from Mongo and providers."""
import asyncio
from copy import deepcopy
from unittest.mock import AsyncMock, MagicMock
import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from pymongo.errors import DuplicateKeyError
from test_monthly_close import inputs, event, snapshot, NOW
from contractor_ops import monthly_close as mc, monthly_close_router as monthly
from contractor_ops import contract_items_router as routes
from contractor_ops.contract_items import attach_contract_items, validate_item, ensure_indexes
from contractor_ops.contract_items_export import QUANTITY_HEADERS
from contractor_ops.monthly_close_export import build_monthly_close_xlsx


def item(**overrides):
    return dict(dict(id='i1', project_id='p1', company_id='c1', source='stage', stage_id='s1',
        basis='unit', code='01', description='צנרת', unit='מ״א', qty=38, qty_by_unit_type=None,
        unit_price=None, price_by_unit_type=None, active=True, measurements={}), **overrides)


def enriched(ins=None, items=None):
    ins = ins or inputs([event('u1'), event('u2')])
    return attach_contract_items(mc.build_account('2026-09', **ins),
                                 {**ins, 'contract_items': [item()] if items is None else items})


def line(result):
    return result['contractors'][0]['items'][0]


def previous(qty=38, with_line=True):
    snap = snapshot(keys=[('s1', 'u3')])
    if with_line:
        snap['contractors'][0]['items'] = [{'item_id': 'i1', 'cumulative_qty': qty, 'cumulative_amount': None}]
    return snap


def test_uniform_types_and_no_items():
    account = enriched()
    row = line(account)
    assert (row['prev_cumulative_qty'], row['this_month_qty'], row['cumulative_qty']) == (0, 76, 76)
    assert (row['units_this_month'], row['evidence_rows'], row['corrections']) == (2, [1, 2], 0)
    assert account['contractors'][0]['has_items'] and account['contractors'][0]['totals_money'] is None
    synthetic = account['contractors'][0]['items'][1]
    assert (synthetic['source'], synthetic['unit'], synthetic['description']) == ('stage_units', 'דירות', 'חשמל (שלב בלי סעיף)')
    assert synthetic['this_month_qty'] == synthetic['cumulative_qty'] - synthetic['prev_cumulative_qty']
    other = account['contractors'][1]
    assert (other['items'], other['has_items'], other['totals_money'], other['account_no']) == ([], False, None, 1)
    ins = inputs([event('u1'), event('u2')])
    ins['units'][0]['unit_type_tag'] = '3 חדרים'
    assert line(enriched(ins, [item(qty_by_unit_type={'3 חדרים': 34})]))['cumulative_qty'] == 72
    assert not enriched(items=[item(active=False)])['contractors'][0]['items']


@pytest.mark.parametrize('new,still,with_line,expected', [
    (2, True, True, (38, 76, 114, 0)), (1, False, True, (38, 0, 38, 1)),
    (0, False, True, (38, -38, 0, 1)), (1, True, False, (38, 38, 76, 0)),
])
def test_previous_corrections_and_new_items(new, still, with_line, expected):
    cells = [event(f'u{i}') for i in range(1, new + 1)]
    if still:
        cells.append(event('u3', date='2026-08-06T10:00:00+00:00'))
    row = line(enriched(inputs(cells, [previous(with_line=with_line)])))
    assert tuple(row[k] for k in ('prev_cumulative_qty', 'this_month_qty', 'cumulative_qty', 'corrections')) == expected


def test_repricing_and_prices():
    snap = snapshot(keys=[('s1', 'u1'), ('s1', 'u2')])
    snap['contractors'][0]['items'] = [{'item_id': 'i1', 'cumulative_qty': 76, 'cumulative_amount': None}]
    ins = inputs([event('u1', date='2026-08-06'), event('u2', date='2026-08-06')], [snap])
    result = enriched(ins, [item(qty=40)])
    assert (line(result)['cumulative_qty'], line(result)['this_month_qty']) == (80, 4)
    assert result['contractors'][0]['stages'][0]['still_unit_ids'] == ['u1', 'u2']
    assert all('still_unit_ids' in s for c in result['contractors'] for s in c['stages'])
    ins = inputs([event('u1'), event('u2')])
    ins['units'][0]['unit_type_tag'] = '5 חדרים'
    priced = enriched(ins, [item(qty=1, unit_price=4000, price_by_unit_type={'5 חדרים': 5000})])
    assert line(priced)['cumulative_amount'] == line(priced)['this_month_amount'] == 9000
    assert priced['contractors'][0]['totals_money'] == {'this_month': 9000, 'cumulative': 9000}


def test_measurements_numbers_snapshots_and_scope():
    measured = item(source='measured', stage_id=None, measurements={'2026-08': {'qty': 82}, '2026-09': {'qty': 36}})
    result = enriched(items=[measured])
    row = result['contractors'][0]['items'][-1]
    assert (row['prev_cumulative_qty'], row['this_month_qty'], row['cumulative_qty'], row['measurement']) == (82, 36, 118, {'qty': 36})
    del measured['measurements']['2026-09']
    row = enriched(items=[measured])['contractors'][0]['items'][-1]
    assert (row['this_month_qty'], row['measurement']) == (0, None)
    ins = inputs(snapshots=[snapshot(month='2026-07'), snapshot(month='2026-08')])
    result = enriched(ins)
    assert [c['account_no'] for c in result['contractors']] == [3, 1]
    frozen = mc.build_snapshot(result, {'id': 'p1'}, {'id': 'pm'}, '', NOW)
    for key in ('items', 'has_items', 'totals_money', 'account_no'):
        assert frozen['contractors'][0][key] == result['contractors'][0][key]
    scoped = monthly._scope_account(frozen, {'role': 'contractor', 'company_id': 'c1', 'can_write': False})
    assert len(scoped['contractors']) == 1
    result['contractors'][0]['items'].clear()
    assert frozen['contractors'][0]['items']


def test_validation_and_indexes():
    args = dict(stages={'s1'}, company_ids={'c1'}, unit_type_tags={'3 חדרים'})
    clean = validate_item({**item(), 'code': '  ', 'description': ' צנרת '}, **args)
    assert clean['code'] is None and clean['description'] == 'צנרת'
    for change in ({'qty': 0}, {'qty': True}, {'qty': float('nan')}, {'unit_price': -1},
                   {'basis': 'floor'}, {'stage_id': 'missing'}, {'company_id': 'missing'},
                   {'qty_by_unit_type': {'missing': 3}}, {'price_by_unit_type': {'3 חדרים': 3}}):
        with pytest.raises(HTTPException) as exc:
            validate_item({**item(), **change}, **args)
        assert exc.value.status_code == 422
    with pytest.raises(HTTPException):
        validate_item(item(source='measured', stage_id=None), existing=item(), **args)
    db = MagicMock()
    db.contract_items.create_index = AsyncMock()
    asyncio.run(ensure_indexes(db))
    assert db.contract_items.create_index.await_args_list[-1].kwargs['partialFilterExpression'] == {'code': {'$type': 'string'}, 'active': True}


@pytest.fixture
def api(monkeypatch):
    db = MagicMock()
    db.contract_items.find_one = AsyncMock(return_value=None)
    db.contract_items.insert_one = AsyncMock()
    db.contract_items.update_one = AsyncMock(return_value=MagicMock(matched_count=1))
    db.monthly_progress_closes.find_one = AsyncMock(return_value=None)
    access = AsyncMock(return_value={'role': 'project_manager', 'can_write': True, 'company_id': None})
    monkeypatch.setattr(routes, 'get_db', lambda: db)
    monkeypatch.setattr(routes, '_access', access)
    monkeypatch.setattr(routes, '_audit', AsyncMock())
    monkeypatch.setattr(routes, 'current_il_month', lambda: '2026-09')
    monkeypatch.setattr(routes, 'load_inputs', AsyncMock(return_value={**inputs(), 'contract_items': [item(), item(id='other', company_id='c2')]}))
    return db, access


def test_create_permissions_duplicate_validation_and_get(api):
    db, access = api
    user = {'id': 'pm'}
    db.contract_items.find_one.return_value = item()
    with pytest.raises(HTTPException) as exc:
        asyncio.run(routes.create_item('p1', item(), user))
    assert exc.value.status_code == 409
    db.contract_items.find_one.return_value = None
    db.contract_items.insert_one.side_effect = DuplicateKeyError('duplicate')
    with pytest.raises(HTTPException) as exc:
        asyncio.run(routes.create_item('p1', item(), user))
    assert exc.value.status_code == 409
    db.contract_items.insert_one.side_effect = None
    for change in ({'basis': 'floor'}, {'stage_id': 'unknown'}):
        with pytest.raises(HTTPException) as exc:
            asyncio.run(routes.create_item('p1', {**item(), **change}, user))
        assert exc.value.status_code == 422
    created = asyncio.run(routes.create_item('p1', item(), user))
    assert created['active'] and created['created_by'] == 'pm' and created['measurements'] == {}
    access.return_value = {'role': 'management_team', 'can_write': False}
    with pytest.raises(HTTPException) as exc:
        asyncio.run(routes.create_item('p1', item(), user))
    assert exc.value.status_code == 403
    access.return_value = {'role': 'contractor', 'can_write': False, 'company_id': 'c1'}
    assert len(asyncio.run(routes.list_items('p1', user))['items']) == 1


def test_update_and_deactivate(api):
    db, _ = api
    db.contract_items.find_one.side_effect = [item(), None]
    updated = asyncio.run(routes.update_item('p1', 'i1', {'description': 'חדש'}, {'id': 'pm'}))
    assert updated['description'] == 'חדש' and updated['updated_by'] == 'pm'
    db.contract_items.find_one.side_effect = None
    db.contract_items.find_one.return_value = item()
    deactivated = asyncio.run(routes.update_item('p1', 'i1', {'active': False}, {'id': 'pm'}))
    assert not deactivated['active'] and deactivated['deactivated_at']
    assert deactivated['description'] == item()['description']
    db.contract_items.find_one.return_value = None
    with pytest.raises(HTTPException) as exc:
        asyncio.run(routes.update_item('p1', 'i1', {}, {'id': 'pm'}))
    assert exc.value.status_code == 404


def test_measurement_guards_and_save(api):
    db, _ = api
    measured = item(source='measured', stage_id=None)
    for doc, closed, month, status in [(measured, {}, '2026-09', None),
            (measured, {'id': 'close'}, '2026-09', 409), (item(), None, '2026-09', 422),
            (measured, None, '2026-10', 422), (item(source='measured', active=False), None, '2026-09', 422)]:
        db.contract_items.find_one.return_value, db.monthly_progress_closes.find_one.return_value = doc, closed
        if status:
            with pytest.raises(HTTPException) as exc:
                asyncio.run(routes.set_measurement('p1', 'i1', month, {'qty': 36}, {'id': 'pm'}))
            assert exc.value.status_code == status
        else:
            saved = asyncio.run(routes.set_measurement('p1', 'i1', month, {'qty': 36, 'note': 'קיר'}, {'id': 'pm', 'name': 'מנהל'}))
            entry = saved['measurements'][month]
            assert entry['qty'] == 36 and entry['by'] == {'id': 'pm', 'name': 'מנהל'} and entry['at']
            assert db.contract_items.update_one.await_args.args[1]['$set'][f'measurements.{month}'] == entry
    db.contract_items.find_one.return_value, db.monthly_progress_closes.find_one.return_value = measured, None
    asyncio.run(routes.set_measurement('p1', 'i1', '2026-09', {'qty': None}, {'id': 'pm'}))
    assert db.contract_items.update_one.await_args.args[1] == {'$unset': {'measurements.2026-09': ''}}


@pytest.mark.parametrize('price', [None, 0, 4000])
def test_export_sheet_headers_numbers_evidence_and_legacy(price):
    account = enriched(items=[item(unit_price=price)])
    contractor = account['contractors'][0]
    ws = load_workbook(build_monthly_close_xlsx({'id': 'p1'}, account, contractor, ''), data_only=False)
    assert ws.sheetnames[0] == 'כמויות לחשבון'
    sheet = ws.worksheets[0]
    assert [c.value for c in sheet[4]] == QUANTITY_HEADERS
    assert 'כמויות לחשבון 1' in sheet['A1'].value
    assert sheet['D5'].data_type == sheet['E5'].data_type == sheet['F5'].data_type == 'n'
    assert sheet['K5'].value == 'גיליון ראיות · שורות 1–2'
    assert any(r[0] == 'סה״כ' for r in sheet.values) == (price is not None)
    legacy = deepcopy(contractor)
    legacy.pop('items')
    assert load_workbook(build_monthly_close_xlsx({'id': 'p1'}, account, legacy, '')).sheetnames == ['סיכום', 'ראיות']
