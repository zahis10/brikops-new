"""Batch 602 contract tests. All data is in-memory; no cell or database writes."""
import asyncio
from datetime import datetime, timezone
from unittest.mock import AsyncMock, MagicMock

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from contractor_ops import monthly_close as mc
from contractor_ops import monthly_close_router as routes
from contractor_ops.monthly_close_export import build_monthly_close_xlsx


NOW = datetime(2026, 9, 30, 12, tzinfo=timezone.utc)


def event(unit, stage='s1', date='2026-09-06T10:00:00+00:00', source='manual',
          status='completed', audit=True):
    cell = {'unit_id': unit, 'stage_id': stage, 'status': status,
            'last_updated_at': date, 'synced_from_qc': source == 'qc'}
    if audit:
        cell['audit'] = [{
            'status_before': 'in_progress', 'status_after': 'completed',
            'timestamp': date, 'actor_name': 'יוסי',
            **({'source': 'qc_sync'} if source == 'qc' else {}),
        }]
    return cell


def snapshot(month='2026-08', keys=(), corrections=(), baseline=(), company='c1'):
    stages = [{'stage_id': stage, 'unit_ids_this_month': [unit]} for stage, unit in keys]
    return {
        'month': month, 'baseline_keys': [list(pair) for pair in baseline],
        'contractors': [{'company_id': company, 'stages': stages}],
        'corrections': [
            {'stage_id': stage, 'unit_id': unit} for stage, unit in corrections
        ],
    }


def inputs(cells=None, snapshots=None, mapping=None, now=NOW):
    buildings = [
        {'id': 'b1', 'name': 'מגדל א'}, {'id': 'b2', 'name': 'מגדל ב'},
    ]
    floors = [
        {'id': 'f1', 'building_id': 'b1', 'floor_number': 1},
        {'id': 'f2', 'building_id': 'b2', 'floor_number': 2},
    ]
    units = [
        {'id': f'u{i}', 'unit_no': str(i), 'floor_id': 'f1' if i <= 3 else 'f2',
         'building_id': 'b1' if i <= 3 else 'b2'}
        for i in range(1, 7)
    ]
    return {
        'stages': [
            {'id': 's1', 'title': 'אינסטלציה', 'scope': 'unit', 'order': 1},
            {'id': 's2', 'title': 'חשמל', 'scope': 'floor', 'order': 2},
            {'id': 's3', 'title': 'ריצוף', 'scope': 'unit', 'order': 3},
        ],
        'units': units, 'buildings': buildings, 'floors': floors,
        'cells': cells or [],
        'companies': [
            {'id': 'c1', 'name': 'אלף', 'trade': 'plumbing'},
            {'id': 'c2', 'name': 'בית', 'trade': 'flooring'},
        ],
        'settings': {
            'stage_companies': mapping if mapping is not None else
            {'s1': 'c1', 's2': 'c1', 's3': 'c2'},
            'contractor_visible': False,
        },
        'snapshots': snapshots or [], 'runs_by_floor': {'f1': 'run-floor'},
        'runs_by_unit': {'u1': 'run-unit'}, 'now': now,
    }


def account(cells=None, snapshots=None, mapping=None, month='2026-09', now=NOW):
    return mc.build_account(
        month, **inputs(cells, snapshots, mapping, now),
    )


def stage(result, stage_id='s1'):
    return next(s for c in result['contractors'] for s in c['stages']
                if s['stage_id'] == stage_id)


def company(result, company_id='c1'):
    return next(c for c in result['contractors'] if c['company_id'] == company_id)


@pytest.mark.parametrize('month,start,end', [
    ('2026-09', '2026-08-31T21:00:00+00:00', '2026-09-30T21:00:00+00:00'),
    ('2026-11', '2026-10-31T22:00:00+00:00', '2026-11-30T22:00:00+00:00'),
])
def test_month_bounds_israel_dst(month, start, end):
    assert mc.month_bounds_utc(month) == (start, end)


@pytest.mark.parametrize('month', ['0001-01', '9999-12'])
def test_extreme_month_bounds_rejected_as_422(month):
    with pytest.raises(HTTPException) as error:
        mc.month_bounds_utc(month)
    assert error.value.status_code == 422
    assert error.value.detail == 'חודש לא תקין (YYYY-MM)'


def test_calendar_wraps_year():
    assert mc.next_month('2026-12') == '2027-01'
    assert mc.prev_month('2027-01') == '2026-12'


def test_calendar_label():
    assert mc.month_label('2026-09') == 'ספטמבר 2026'


def test_current_il_month_at_utc_boundary():
    assert mc.current_il_month(datetime(2026, 8, 31, 22, tzinfo=timezone.utc)) == '2026-09'


def test_completed_manual_uses_transition_not_note_edit():
    cell = event('u1')
    cell['audit'].append({'status_after': 'completed', 'status_before': 'completed',
                          'timestamp': '2026-09-20T00:00:00+00:00'})
    assert mc.completion_event(cell)['completed_at'] == cell['audit'][0]['timestamp']
    assert mc.completion_event(cell)['source'] == 'manual'


def test_qc_transition_source():
    assert mc.completion_event(event('u1', source='qc'))['source'] == 'qc'


def test_recompleted_uses_second_transition():
    cell = event('u1')
    cell['audit'].extend([
        {'status_before': 'completed', 'status_after': 'in_progress',
         'timestamp': '2026-09-10T00:00:00+00:00'},
        {'status_before': 'in_progress', 'status_after': 'completed',
         'timestamp': '2026-09-11T00:00:00+00:00', 'source': 'qc_sync'},
    ])
    assert mc.completion_event(cell)['completed_at'].startswith('2026-09-11')
    assert mc.completion_event(cell)['source'] == 'qc'


def test_partial_never_counts():
    assert mc.completion_event(event('u1', status='partial')) is None


def test_legacy_timestamp_source():
    assert mc.completion_event(event('u1', audit=False, source='qc'))['source'] == 'qc'


def test_legacy_without_timestamp_does_not_count():
    cell = event('u1', audit=False)
    cell.pop('last_updated_at')
    assert mc.completion_event(cell) is None


def test_single_trade_suggestion():
    assert mc.suggest_stage_companies(
        [{'id': 's1', 'title': 'אינסטלציה'}],
        [{'id': 'c1', 'trade': 'plumbing'}],
    ) == {'s1': 'c1'}


def test_ambiguous_trade_no_suggestion():
    assert mc.suggest_stage_companies(
        [{'id': 's1', 'title': 'אינסטלציה'}],
        [{'id': 'c1', 'trade': 'plumbing'}, {'id': 'c2', 'trade': 'plumbing'}],
    ) == {}


def test_custom_trade_label_suggestion():
    assert mc.suggest_stage_companies(
        [{'id': 's1', 'title': 'התקנת פרקט'}],
        [{'id': 'c1', 'trade': 'custom', 'trade_label': 'פרקטים'}],
    ) == {'s1': 'c1'}


@pytest.mark.parametrize('mapping', [{'unknown': 'c1'}, {'s1': 'other'}])
def test_settings_rejects_foreign_stage_or_company(mapping):
    with pytest.raises(HTTPException) as error:
        mc.validate_settings({'stage_companies': mapping}, ['s1'], ['c1'])
    assert error.value.status_code == 422


def test_settings_accepts_null_unmapping():
    assert mc.validate_settings(
        {'stage_companies': {'s1': None}, 'contractor_visible': True},
        ['s1'], ['c1'],
    )['stage_companies'] == {'s1': None}


def test_default_visibility_off():
    assert mc.default_settings()['contractor_visible'] is False


def test_paid_keys_counted_snapshot():
    assert mc.paid_keys([snapshot(keys=[('s1', 'u1')])]) == {('s1', 'u1'): '2026-08'}


def test_paid_keys_baseline():
    assert mc.paid_keys([snapshot(baseline=[('s1', 'u1')])]) == {('s1', 'u1'): mc.BASELINE}


def test_paid_keys_closed_correction_removes_key():
    snaps = [snapshot(keys=[('s1', 'u1')]),
             snapshot('2026-09', corrections=[('s1', 'u1')])]
    assert ('s1', 'u1') not in mc.paid_keys(snaps)


def test_paid_keys_chronological_recount_after_repayment():
    snaps = [snapshot('2026-10', keys=[('s1', 'u1')]),
             snapshot('2026-09', corrections=[('s1', 'u1')]),
             snapshot(keys=[('s1', 'u1')])]
    assert mc.paid_keys(snaps)[('s1', 'u1')] == '2026-10'


def test_first_account_only_window_counts():
    cells = [event(f'u{i}', date='2026-09-0%dT10:00:00+00:00' % i,
                   source='qc' if i <= 2 else 'manual') for i in range(1, 5)]
    cells += [event('u5', date='2026-08-25T10:00:00+00:00'),
              event('u6', date='2026-08-26T10:00:00+00:00')]
    result = account(cells)
    assert result['kpis']['this_month'] == 4
    assert result['kpis']['qc'] == 2
    assert result['kpis']['manual'] == 2
    assert stage(result)['cumulative'] == 4
    assert stage(result)['pct'] == 67
    assert result['applying'] is False
    assert [r['unit_no'] for r in stage(result)['evidence']] == ['1', '2', '3', '4']
    assert stage(result)['by_building'][0]['this_month'] == 3
    assert stage(result)['by_building'][1]['this_month'] == 1


def test_unmapped_stage_excluded_and_listed():
    result = account([event('u1', stage='s3')], mapping={'s1': 'c1', 's2': 'c1'})
    assert result['kpis']['this_month'] == 0
    assert result['unmapped_stages'][0]['id'] == 's3'


def test_no_first_month_late_from_august():
    result = account([event('u1', date='2026-08-31T20:00:00+00:00')])
    assert result['kpis']['this_month'] == 0


def test_paid_late_addition_after_prior_close():
    result = account([event('u1', date='2026-08-29T12:00:00+00:00')],
                     [snapshot(keys=[('s1', 'u2')])])
    assert stage(result)['this_month'] == 1
    assert stage(result)['evidence'][0]['late_from_month'] == '2026-08'
    assert stage(result)['cumulative'] == 1


def test_reopened_six_paid_results_in_five_not_four():
    prior = snapshot(keys=[('s1', f'u{i}') for i in range(1, 7)])
    cells = [event(f'u{i}', date='2026-08-10T00:00:00+00:00') for i in range(1, 6)]
    cells.append(event('u6', status='in_progress'))
    result = account(cells, [prior])
    assert result['kpis']['corrections'] == 1
    assert stage(result)['cumulative'] == 5
    assert company(result)['corrections'][0]['counted_in_month'] == '2026-08'


def test_recompleted_paid_key_has_no_second_addition_or_correction():
    prior = snapshot(keys=[('s1', f'u{i}') for i in range(1, 7)])
    cells = [event(f'u{i}', date='2026-08-10T00:00:00+00:00') for i in range(1, 6)]
    cells.append(event('u6', date='2026-09-11T00:00:00+00:00'))
    result = account(cells, [prior])
    assert stage(result)['this_month'] == 0
    assert stage(result)['still_paid'] == 6
    assert stage(result)['cumulative'] == 6
    assert not company(result)['corrections']


def test_invariant_with_late_addition_and_reopened_paid_key():
    prior = snapshot(keys=[('s1', f'u{i}') for i in range(1, 5)])
    cells = [event(f'u{i}', date='2026-08-10T00:00:00+00:00') for i in range(1, 4)]
    cells += [event('u4', status='in_progress'),
              event('u5', date='2026-08-29T00:00:00+00:00')]
    result = account(cells, [prior])
    assert stage(result)['cumulative'] == 4 - 1 + 1


def test_first_snapshot_stores_prior_work_as_baseline():
    result = account([event('u1', date='2026-08-10T00:00:00+00:00'),
                      event('u2')])
    doc = mc.build_snapshot(result, {'id': 'p1'}, {'id': 'pm', 'name': 'מנהל'},
                            '  מידע  ', NOW)
    assert ['s1', 'u1'] in doc['baseline_keys']
    assert ['s1', 'u2'] not in doc['baseline_keys']
    assert doc['closed_by'] == {'id': 'pm', 'name': 'מנהל'}
    assert doc['note'] == 'מידע'


@pytest.mark.parametrize('utc_stamp,local_day', [
    ('2026-09-30T20:59:59+00:00', '2026-09-30'),
    ('2026-09-30T21:00:00+00:00', '2026-10-01'),
])
def test_snapshot_closed_date_uses_israel_day_not_utc_day(utc_stamp, local_day):
    instant = datetime.fromisoformat(utc_stamp)
    result = account([event('u1')])
    doc = mc.build_snapshot(result, {'id': 'p1'}, {'id': 'pm', 'name': 'מנהל'},
                            None, instant)
    assert doc['closed_at'] == utc_stamp
    assert doc['closed_date_il'] == local_day


def test_baseline_reopen_is_correction():
    prior = snapshot(baseline=[('s1', 'u1')])
    result = account([event('u1', status='in_progress')], [prior])
    assert company(result)['corrections'][0]['counted_in_month'] == mc.BASELINE


def test_second_snapshot_has_no_baseline_and_flat_corrections():
    prior = snapshot(keys=[('s1', 'u1')])
    result = account([event('u1', status='in_progress')], [prior])
    doc = mc.build_snapshot(result, {'id': 'p1'}, {'id': 'pm', 'name': 'מנהל'},
                            'n' * 600, NOW)
    assert doc['baseline_keys'] == []
    assert len(doc['note']) == 500
    assert doc['corrections'][0]['stage_id'] == 's1'
    assert doc['corrections'][0]['unit_id'] == 'u1'


def test_unmapping_paid_stage_preserves_old_company_correction():
    prior = snapshot(keys=[('s3', 'u1')], company='c2')
    result = account([event('u1', stage='s3', status='in_progress')],
                     [prior], mapping={'s1': 'c1'})
    assert company(result, 'c2')['corrections'][0]['unit_id'] == 'u1'
    assert company(result, 'c2')['stages'] == []


def test_remapping_paid_stage_correction_belongs_to_historical_company():
    prior = snapshot(keys=[('s1', 'u1')], company='c1')
    result = account([event('u1', status='in_progress')],
                     [prior], mapping={'s1': 'c2'})
    assert company(result, 'c1')['corrections'][0]['company_id'] == 'c1'


def test_paid_key_reenters_only_after_closed_correction():
    prior = [snapshot(keys=[('s1', 'u1')]),
             snapshot('2026-09', corrections=[('s1', 'u1')])]
    result = account([event('u1', date='2026-10-10T00:00:00+00:00')],
                     prior, month='2026-10', now=datetime(2026, 10, 31, tzinfo=timezone.utc))
    assert stage(result)['this_month'] == 1


@pytest.mark.parametrize('month,expected', [
    ('2026-08', False), ('2026-09', True), ('2026-10', False),
])
def test_only_next_nonfuture_month_closable(month, expected):
    assert account(snapshots=[snapshot()], month=month)['closable'] is expected


def test_first_close_any_past_month_eligible():
    assert account(month='2026-07')['closable'] is True


def test_malformed_completion_timestamp_is_skipped():
    assert account([event('u1', date='not-a-date')])['kpis']['this_month'] == 0


def export_fixture(corrections=False, unit_no='=1+1'):
    row = {'unit_id': 'u1', 'unit_no': unit_no, 'building_name': 'מגדל א',
           'floor_number': 1, 'floor_id': 'f1', 'stage_id': 's1',
           'completed_at': '2026-09-04T21:30:00+00:00', 'actor_name': 'יוסי',
           'source': 'qc', 'qc_run_id': 'run1', 'late_from_month': None}
    correction = {'stage_id': 's1', 'stage_title': 'ריצוף',
                  'unit_id': 'u2', 'unit_no': '=1+1',
                  'building_name': 'מגדל א', 'counted_in_month': '2026-08'}
    contractor = {'company_id': 'c1', 'name': '=נוסחה',
                  'stages': [{'stage_id': 's1', 'title': 'ריצוף',
                              'by_building': [{'building_id': 'b1',
                                               'building_name': 'מגדל א',
                                               'total_units': 3, 'this_month': 1,
                                               'cumulative': 2, 'qc_count': 1,
                                               'manual_count': 0}],
                              'evidence': [row]}],
                  'totals': {'this_month': 1, 'cumulative': 2, 'qc': 1, 'manual': 0},
                  'corrections': [correction] if corrections else []}
    doc = {'id': 'p1', 'name': 'הפרויקט'}
    account_doc = {'month': '2026-09', 'label': 'ספטמבר 2026',
                   'status': 'open'}
    return load_workbook(build_monthly_close_xlsx(
        doc, account_doc, contractor, 'https://app.example',
    ))


def test_export_exact_headers_and_rtl():
    wb = export_fixture()
    assert wb.sheetnames == ['סיכום', 'ראיות']
    assert [c.value for c in wb['סיכום'][4]] == [
        'קבלן', 'שלב', 'בניין', 'סה״כ דירות', 'בוצע החודש', 'מצטבר',
        'אחוז', 'אושרו בבקרת ביצוע', 'סומנו ידנית',
    ]
    assert [c.value for c in wb['ראיות'][1]] == [
        '#', 'בניין', 'קומה', 'דירה', 'שלב', 'תאריך ביצוע',
        'אושר ע״י', 'מקור', 'קישור',
    ]
    assert all(sheet.sheet_view.rightToLeft for sheet in wb)
    assert wb['סיכום'].freeze_panes == 'A5'


def test_export_evidence_qc_link_date_and_nonformula():
    wb = export_fixture()
    row = wb['ראיות'][2]
    assert row[3].value == '=1+1'
    assert row[3].data_type == 's'
    assert row[5].value == '5.9.2026'
    assert row[8].value == (
        'https://app.example/projects/p1/floors/f1/qc/run1/stage/s1?unitId=u1'
    )
    assert wb['סיכום']['A1'].data_type == 's'


def test_export_optional_corrections_and_negative_row():
    wb = export_fixture(corrections=True)
    assert wb.sheetnames == ['סיכום', 'ראיות', 'תיקונים']
    assert wb['תיקונים']['C2'].data_type == 's'
    assert wb['תיקונים']['D2'].value == 'אוגוסט 2026'
    assert wb['סיכום']['E7'].value == -1


def test_export_baseline_and_malformed_date():
    doc = {'id': 'p1'}
    contractor = {'name': 'קבלן', 'stages': [{'title': 'צבע', 'evidence': [
        {'unit_no': '1', 'completed_at': 'broken', 'source': 'manual'},
    ]}], 'corrections': [
        {'stage_title': 'צבע', 'counted_in_month': 'baseline', 'unit_no': '2'},
    ]}
    wb = load_workbook(build_monthly_close_xlsx(
        doc, {'status': 'closed', 'closed_at': 'bad', 'closed_by': {'name': 'מנהלת'}},
        contractor, '',
    ))
    assert wb['ראיות']['F2'].value == '—'
    assert wb['ראיות']['I2'].value == '/projects/p1/execution-matrix'
    assert 'לפני החשבון הראשון ב-BrikOps' in wb['סיכום']['I6'].value


@pytest.mark.parametrize('source,expected', [
    ('qc', 'אושר בבקרת ביצוע · אחרי סגירת אוגוסט 2026'),
    ('manual', 'סומן ידנית · אחרי סגירת אוגוסט 2026'),
])
def test_export_late_source_and_links(source, expected):
    row = {'unit_id': 'u1', 'unit_no': '1', 'stage_id': 's1',
           'floor_id': 'f1', 'qc_run_id': 'run1', 'source': source,
           'late_from_month': '2026-08'}
    contractor = {'name': 'קבלן', 'stages': [{'title': 'חשמל',
                                            'by_building': [], 'evidence': [row]}]}
    wb = load_workbook(build_monthly_close_xlsx(
        {'id': 'p1'}, {'status': 'open', 'label': 'ספטמבר 2026'},
        contractor, '',
    ))
    assert wb['ראיות']['H2'].value == expected
    if source == 'qc':
        assert wb['ראיות']['I2'].value == (
            '/projects/p1/floors/f1/qc/run1/stage/s1?unitId=u1'
        )
    else:
        assert wb['ראיות']['I2'].value == '/projects/p1/execution-matrix'


def test_export_closed_at_uses_israel_date():
    wb = load_workbook(build_monthly_close_xlsx(
        {'id': 'p1'}, {
            'status': 'closed', 'closed_at': '2026-09-30T22:00:00+00:00',
            'closed_by': {'name': 'רונית'},
        }, {'name': 'קבלן'}, '',
    ))
    assert wb['סיכום']['A2'].value == 'נסגר 1.10.2026 ע״י רונית'


def test_contractor_scope_never_leaks_snapshot_fields():
    doc = account([event('u1'), event('u2', 's3')])
    doc.update({'status': 'closed', 'is_snapshot': True,
                'note': 'private note: other contractor',
                'corrections': [
                    {'company_id': 'c2', 'stage_id': 's3', 'unit_id': 'u2'},
                ], 'baseline_keys': [['s1', 'u1'], ['s3', 'u2']]})
    scoped = routes._scope_account(
        doc, {'role': 'contractor', 'company_id': 'c1', 'can_write': False},
    )
    assert [c['company_id'] for c in scoped['contractors']] == ['c1']
    assert scoped['kpis']['this_month'] == 1
    assert scoped['kpis']['corrections'] == 0
    assert scoped['corrections'] == []
    assert scoped['baseline_keys'] == [['s1', 'u1']]
    assert scoped['settings_snapshot']['stage_companies'] == {'s1': 'c1', 's2': 'c1'}
    assert scoped['unmapped_stages'] == []
    assert 'note' not in scoped
    assert scoped['permissions']['can_close'] is False
    assert len(doc['contractors']) == 2  # did not mutate the manager's object


def test_contractor_scope_open_unmapped_and_kpis():
    doc = account([event('u1'), event('u2', stage='s3')])
    doc['unmapped_stages'] = [{'id': 's4', 'title': 'private'}]
    scoped = routes._scope_account(
        doc, {'role': 'contractor', 'company_id': 'c2', 'can_write': False},
    )
    assert scoped['kpis']['this_month'] == 1
    assert scoped['totals']['stages'] == 1
    assert scoped['unmapped_stages'] == []


def mock_router_db(monkeypatch, project=None):
    db = MagicMock()
    db.projects.find_one = AsyncMock(return_value=project or {'id': 'p1', 'name': 'פרויקט'})
    db.project_memberships.find_one = AsyncMock(return_value=None)
    db.monthly_progress_closes.find_one = AsyncMock(return_value=None)
    db.monthly_progress_closes.insert_one = AsyncMock()
    db.projects.update_one = AsyncMock(return_value=MagicMock(matched_count=1))
    monkeypatch.setattr(routes, 'get_db', lambda: db)
    monkeypatch.setattr(routes, '_is_super_admin', lambda user: False)
    monkeypatch.setattr(routes, '_get_project_role',
                        AsyncMock(return_value='project_manager'))
    fixture_inputs = inputs()
    fixture_inputs.pop('now')
    monkeypatch.setattr(routes, 'load_inputs',
                        AsyncMock(return_value={**fixture_inputs, 'project': {'id': 'p1', 'name': 'פרויקט'}}))
    monkeypatch.setattr(routes, '_now', lambda: NOW)
    monkeypatch.setattr(routes, 'current_il_month', lambda: '2026-09')
    monkeypatch.setattr(routes, '_audit', AsyncMock())
    return db


def test_access_opted_out_contractor_forbidden(monkeypatch):
    db = mock_router_db(monkeypatch)
    monkeypatch.setattr(routes, '_get_project_role', AsyncMock(return_value='contractor'))
    db.project_memberships.find_one = AsyncMock(return_value={'company_id': 'c1'})
    with pytest.raises(HTTPException) as error:
        asyncio.run(routes.get_month('p1', '2026-09', {'id': 'u1'}))
    assert error.value.status_code == 403


def test_access_contractor_requires_company(monkeypatch):
    db = mock_router_db(monkeypatch, project={
        'id': 'p1', 'monthly_close_settings': {'contractor_visible': True},
    })
    monkeypatch.setattr(routes, '_get_project_role', AsyncMock(return_value='contractor'))
    db.project_memberships.find_one = AsyncMock(return_value={})
    with pytest.raises(HTTPException) as error:
        asyncio.run(routes._access({'id': 'u1'}, 'p1', db))
    assert error.value.status_code == 403


def test_contractor_opted_in_get_is_scoped_read_only(monkeypatch):
    db = mock_router_db(monkeypatch, project={
        'id': 'p1', 'monthly_close_settings': {'contractor_visible': True},
    })
    monkeypatch.setattr(routes, '_get_project_role', AsyncMock(return_value='contractor'))
    db.project_memberships.find_one = AsyncMock(return_value={'company_id': 'c1'})
    monkeypatch.setattr(routes, '_account', AsyncMock(
        return_value=account([event('u1'), event('u2', stage='s3')]),
    ))
    result = asyncio.run(routes.get_month('p1', '2026-09', {'id': 'u1'}))
    assert result['permissions']['can_close'] is False
    assert result['permissions']['scoped_company_id'] == 'c1'
    assert len(result['contractors']) == 1
    assert result['kpis']['this_month'] == 1


@pytest.mark.parametrize('role,write', [('management_team', False),
                                         ('project_manager', True)])
def test_management_get_month_permissions(monkeypatch, role, write):
    mock_router_db(monkeypatch)
    monkeypatch.setattr(routes, '_get_project_role', AsyncMock(return_value=role))
    monkeypatch.setattr(routes, '_account', AsyncMock(return_value=account()))
    result = asyncio.run(routes.get_month('p1', '2026-09', {'id': 'u1'}))
    assert result['permissions']['can_close'] is write
    assert result['permissions']['can_write'] is write


def test_role_checked_before_malformed_month(monkeypatch):
    mock_router_db(monkeypatch)
    monkeypatch.setattr(routes, '_get_project_role', AsyncMock(return_value='contractor'))
    with pytest.raises(HTTPException) as error:
        asyncio.run(routes.get_month('p1', '2026-13', {'id': 'u1'}))
    assert error.value.status_code == 403


def test_authorized_malformed_month_422(monkeypatch):
    mock_router_db(monkeypatch)
    with pytest.raises(HTTPException) as error:
        asyncio.run(routes.get_month('p1', '2026-13', {'id': 'pm'}))
    assert error.value.status_code == 422


def test_management_cannot_put_settings(monkeypatch):
    mock_router_db(monkeypatch)
    monkeypatch.setattr(routes, '_get_project_role',
                        AsyncMock(return_value='management_team'))
    with pytest.raises(HTTPException) as error:
        asyncio.run(routes.put_settings('p1', {}, {'id': 'm'}))
    assert error.value.status_code == 403


def test_settings_optimistic_conflict_does_not_write(monkeypatch):
    db = mock_router_db(monkeypatch)
    project = {'id': 'p1', 'monthly_close_settings': {
        'stage_companies': {}, 'updated_at': 'latest',
    }}
    routes.load_inputs.return_value['project'] = project
    with pytest.raises(HTTPException) as error:
        asyncio.run(routes.put_settings(
            'p1', {'stage_companies': {}, 'updated_at': 'stale'}, {'id': 'pm'},
        ))
    assert error.value.status_code == 409
    db.projects.update_one.assert_not_awaited()


def test_settings_saves_explicit_mapping_and_audits(monkeypatch):
    db = mock_router_db(monkeypatch)
    result = asyncio.run(routes.put_settings('p1', {
        'stage_companies': {'s1': 'c1'}, 'contractor_visible': True,
        'updated_at': None,
    }, {'id': 'pm'}))
    assert result['stage_companies'] == {'s1': 'c1'}
    db.projects.update_one.assert_awaited_once()
    assert db.projects.update_one.await_args.args[1]['$set']['monthly_close_settings'][
        'contractor_visible'] is True
    routes._audit.assert_awaited_once()


def test_close_twice_denied_before_any_insert(monkeypatch):
    db = mock_router_db(monkeypatch)
    db.monthly_progress_closes.find_one.return_value = {'id': 'old'}
    with pytest.raises(HTTPException) as error:
        asyncio.run(routes.close_month('p1', '2026-09', {}, {'id': 'pm'}))
    assert error.value.status_code == 409
    db.monthly_progress_closes.insert_one.assert_not_awaited()


@pytest.mark.parametrize('month', ['2026-10', '2026-08'])
def test_close_future_or_wrong_next_month_denied(monkeypatch, month):
    db = mock_router_db(monkeypatch)
    routes.load_inputs.return_value['snapshots'] = [snapshot()]
    with pytest.raises(HTTPException) as error:
        asyncio.run(routes.close_month('p1', month, {}, {'id': 'pm'}))
    assert error.value.status_code == 422
    db.monthly_progress_closes.insert_one.assert_not_awaited()


def test_close_opt_in_notifies_only_mapped_contractor_members(monkeypatch):
    db = mock_router_db(monkeypatch)
    routes.load_inputs.return_value['settings']['contractor_visible'] = True
    routes.load_inputs.return_value['settings']['stage_companies'] = {'s1': 'c1'}
    db.project_memberships.find.return_value.to_list = AsyncMock(
        return_value=[{'user_id': 'contractor-user'}],
    )
    notify = AsyncMock()
    monkeypatch.setattr(routes, 'create_defect_notification', notify)
    result = asyncio.run(routes.close_month('p1', '2026-09', {},
                                             {'id': 'pm', 'name': 'מנהל'}))
    assert result['status'] == 'closed'
    db.monthly_progress_closes.insert_one.assert_awaited_once()
    notify.assert_awaited_once()
    assert notify.await_args.args[1] == ['contractor-user']
    assert notify.await_args.kwargs['notification_type'] == 'monthly_close'
    assert notify.await_args.kwargs['extra'] == {'month': '2026-09', 'company_id': 'c1'}


def test_export_contractor_cannot_request_other_company(monkeypatch):
    db = mock_router_db(monkeypatch, project={
        'id': 'p1', 'monthly_close_settings': {'contractor_visible': True},
    })
    monkeypatch.setattr(routes, '_get_project_role', AsyncMock(return_value='contractor'))
    db.project_memberships.find_one = AsyncMock(return_value={'company_id': 'c1'})
    with pytest.raises(HTTPException) as error:
        asyncio.run(routes.export_xlsx('p1', '2026-09',
                                       {'company_id': 'c2'}, {'id': 'u1'}))
    assert error.value.status_code == 403


def test_export_nonexistent_contractor_404(monkeypatch):
    mock_router_db(monkeypatch)
    monkeypatch.setattr(routes, '_account', AsyncMock(return_value=account()))
    with pytest.raises(HTTPException) as error:
        asyncio.run(routes.export_xlsx('p1', '2026-09',
                                       {'company_id': 'missing'}, {'id': 'pm'}))
    assert error.value.status_code == 404