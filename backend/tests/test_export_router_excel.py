import asyncio
from io import BytesIO
from unittest.mock import AsyncMock, MagicMock, patch

from openpyxl import load_workbook

from contractor_ops import export_router


class _Cursor:
    def __init__(self, values):
        self._values = values

    async def to_list(self, _limit):
        return list(self._values)


def _formula_cells(workbook):
    return [
        f'{sheet.title}!{cell.coordinate}'
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.data_type == 'f'
    ]


def test_generate_excel_neutralizes_defect_and_map_text():
    output = export_router._generate_excel(
        [{
            'id': 'task-1',
            'display_number': '1',
            'title': '=1+1',
            'description': '@x',
            'status': 'open',
            'building_id': 'building-1',
            'company_id': 'company-1',
        }],
        '+Project',
        {},
        {'company-1': '-Company'},
        {},
        {},
        {'building-1': '=Building'},
    )

    sheet = load_workbook(output).active
    expected = {
        2: '+Project',
        3: '=Building',
        7: '=1+1',
        8: '@x',
        10: '-Company',
    }
    for column, value in expected.items():
        cell = sheet.cell(row=2, column=column)
        assert cell.value == value
        assert cell.data_type == 's'


def test_generate_full_excel_neutralizes_every_write_row_sheet():
    db = MagicMock()
    db.buildings.find.return_value = _Cursor([{
        'id': 'building-1',
        'project_id': 'project-1',
        'name': '=Building',
    }])
    db.floors.find.return_value = _Cursor([])
    db.units.find.return_value = _Cursor([])
    db.project_companies.find.return_value = _Cursor([{
        'id': 'company-1',
        'project_id': 'project-1',
        'name': '@Company',
        'contact_name': '+Contact',
    }])
    db.project_memberships.find.return_value = _Cursor([{
        'user_id': 'user-1',
        'role': 'project_manager',
    }])
    db.users.find.return_value = _Cursor([{
        'id': 'user-1',
        'name': '=Worker',
    }])
    db.tasks.find.return_value = _Cursor([{
        'id': 'task-1',
        'title': '=1+1',
        'description': '@Description',
        'status': 'open',
        'building_id': 'building-1',
    }])
    db.handover_protocols.find.return_value = _Cursor([{
        'id': 'protocol-1',
        'building_id': 'building-1',
        'notes': '-Protocol note',
        'sections': [],
    }])
    db.qc_runs.find.return_value = _Cursor([])

    async def build():
        with (
            patch.object(export_router, 'get_db', return_value=db),
            patch.object(
                export_router,
                '_resolve_image_links',
                new=AsyncMock(return_value=[]),
            ),
        ):
            return await export_router._generate_full_excel(
                'project-1',
                '+Project',
            )

    workbook = load_workbook(BytesIO(asyncio.run(build()).getvalue()))
    assert _formula_cells(workbook) == []
    assert workbook['ליקויים']['B2'].value == '+Project'
    assert workbook['ליקויים']['G2'].value == '=1+1'
    assert workbook['ליקויים']['H2'].value == '@Description'
    assert workbook['פרוטוקולי מסירה']['L2'].value == '-Protocol note'
    assert workbook['צוות']['A2'].value == '=Worker'
    assert workbook['חברות וקבלנים']['A2'].value == '@Company'
    assert workbook['חברות וקבלנים']['B2'].value == '+Contact'
    for coordinate in ('B2', 'G2', 'H2'):
        assert workbook['ליקויים'][coordinate].data_type == 's'